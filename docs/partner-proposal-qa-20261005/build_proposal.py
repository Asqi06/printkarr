from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).parent
OUT = ROOT.parent / 'PrintKarr_Shop_Partner_Proposal_Hindi_2026-10-05.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.65), Inches(.65)
sec.left_margin, sec.right_margin = Inches(.75), Inches(.75)
sec.header_distance, sec.footer_distance = Inches(.25), Inches(.28)

def font(style, size, bold=False):
    style.font.name = 'Nirmala UI'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rp = style.element.get_or_add_rPr()
    rf = rp.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rp.append(rf)
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        rf.set(qn('w:'+key), 'Nirmala UI')
    lang = OxmlElement('w:lang'); lang.set(qn('w:val'),'hi-IN'); lang.set(qn('w:bidi'),'hi-IN'); rp.append(lang)
    cs = OxmlElement('w:szCs'); cs.set(qn('w:val'),str(int(size*2))); rp.append(cs)

font(doc.styles['Normal'], 11.5)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
for name, size in [('Title',24),('Subtitle',12),('Heading 1',17),('Heading 2',13)]:
    font(doc.styles[name],size,name!='Subtitle')
    doc.styles[name].paragraph_format.space_after=Pt(9)
    doc.styles[name].paragraph_format.space_before=Pt(11 if name=='Heading 2' else 0)
    doc.styles[name].paragraph_format.keep_with_next=True
doc.styles['Title'].paragraph_format.line_spacing=1.05

def p(text, bold=False, size=None):
    a=doc.add_paragraph()
    r=a.add_run(text); r.bold=bold
    if size:
        r.font.size=Pt(size)
        x=OxmlElement('w:szCs'); x.set(qn('w:val'),str(int(size*2))); r._element.get_or_add_rPr().append(x)
    return a

def h(text): doc.add_heading(text,level=2)
def page(title): doc.add_page_break(); doc.add_heading(title,level=1)
def table(headers, rows, widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    pr.append(borders)
    for index, values in enumerate([headers]+rows):
        row=t.rows[0] if index==0 else t.add_row()
        rp=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');rp.append(no)
        if index==0:rp.append(OxmlElement('w:tblHeader'))
        for i,value in enumerate(values):
            c=row.cells[i];c.width=Inches(widths[i]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                n=OxmlElement('w:'+side);n.set(qn('w:w'),'100');n.set(qn('w:type'),'dxa');m.append(n)
            cp.append(m)
            if index==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF5');cp.append(sh)
            a=c.paragraphs[0];a.paragraph_format.space_after=Pt(2);a.paragraph_format.line_spacing=1.03
            r=a.add_run(str(value));r.bold=index==0;r.font.size=Pt(10.5)
            z=OxmlElement('w:szCs');z.set(qn('w:val'),'21');r._element.get_or_add_rPr().append(z)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

def field(label):
    a=p(label+'  __________________________________________________')
    a.paragraph_format.space_after=Pt(11)

foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('PrintKarr  |  प्रस्तावित पायलट  |  पृष्ठ ');r.font.size=Pt(9)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)

doc.add_paragraph('PrintKarr दुकान साझेदारी प्रस्ताव',style='Title')
doc.add_paragraph('स्थानीय प्रिंट सेवा के लिए 30 दिन का पायलट',style='Subtitle')
p('दुकान मालिकों के लिए चर्चा और लिखित सहमति का प्रस्ताव\nस्थान  चाला और वापी का स्थानीय क्षेत्र\nप्रस्ताव दिनांक  5 अक्टूबर 2026',size=10.5)
doc.add_heading('1 अपनी दुकान से ऑनलाइन प्रिंट ऑर्डर पूरा करें',level=1)
p('हम चाला, वापी की प्रिंट और ज़ेरॉक्स दुकानों के साथ एक छोटा पायलट शुरू करना चाहते हैं। ग्राहक PrintKarr के माध्यम से फाइल और प्रिंट की जरूरत बताएगा। सहमत पार्टनर दुकान अपने परिसर में प्रिंट तैयार करेगी और PrintKarr ग्राहक से समन्वय तथा उपलब्ध डिलीवरी की व्यवस्था करेगा। आपकी मौजूदा दुकान, मशीन और अनुभव इस साझेदारी का आधार होंगे।')
p('हमारी ग्राहक सेवा अभी शुरुआत में है और फिलहाल ग्राहक संख्या शून्य है। इसलिए यह स्थापित मांग या निश्चित कमाई का दावा नहीं है। हमारा प्रस्ताव 2–3 उपयुक्त दुकानों के साथ 30 दिन तक वास्तविक भुगतान वाले ऑर्डर, गुणवत्ता और लागत पर काम करना है। शुरुआत का क्षेत्र चाला के पास एक छोटा समूह होगा; सामान्य डिलीवरी सीमा लगभग 2–3 किमी पर सहमति के बाद तय होगी।')
h('दुकान ग्राहक और प्लेटफॉर्म के लिए लाभ')
p('दुकान को ऑनलाइन पूछताछ और प्रिंट ऑर्डर पाने का एक अतिरिक्त रास्ता मिल सकता है। दुकान केवल वह काम स्वीकार करेगी जिसे सहमत दर, समय और गुणवत्ता पर पूरा कर सकती है। आपके पुराने ग्राहकों पर कोई रोक या अनिवार्य exclusivity नहीं होगी।')
p('ग्राहक को पहले से स्पष्ट कीमत, अपनी फाइल और सेटिंग चुनने की सुविधा तथा उपलब्ध होने पर भुगतान वाली डिलीवरी या सहमत pickup विकल्प मिलेंगे। PrintKarr की आय पूरे हुए प्रिंट काम की सहमत हिस्सेदारी से होगी; हमारा काम ऑर्डर, ग्राहक संवाद और भुगतान व डिलीवरी का समन्वय करना है।')
h('पहले चर्चा फिर शुरुआत')
p('दुकान की क्षमता, वास्तविक लागत और सेवा क्षेत्र समझने के बाद दोनों पक्ष अंतिम वाणिज्यिक विवरण लिखित रूप से भरेंगे। यह प्रस्ताव पूर्ण कानूनी रूप से सत्यापित अनुबंध होने का दावा नहीं करता। संचालक की वास्तविक कानूनी पहचान, कर व्यवस्था और आवश्यक दस्तावेज शुरुआत से पहले पूरे किए जाएंगे।',size=10.5)

page('2 दरें प्रिंट मानक और प्लेटफॉर्म हिस्सेदारी')
p('निम्न दरें प्रस्तावित पायलट दरें हैं। दुकान को स्वीकार करने से पहले जांचना है कि कागज, toner या ink, बिजली, मेहनत और सामान्य खराब प्रिंट की लागत के बाद उसका काम व्यवहार्य है। असहमत दर पर काम स्वीकार करना आवश्यक नहीं है।')
table(['सेवा','ग्राहक दर प्रति प्रिंट साइड','10% हिस्सेदारी','दुकान का हिस्सा'],[
    ['A4 ब्लैक एंड व्हाइट','₹2.00','₹0.20','₹1.80'],
    ['A4 सामान्य रंगीन दस्तावेज','₹5.00','₹0.50','₹4.50'],
],[2.3,1.7,1.4,1.6])
p('तालिका केवल सरल, बिना कर वाले गणना उदाहरण के लिए है। वास्तविक ग्राहक बिल और दुकान payout में लागू करों का व्यवहार अंतिम विवरण में अलग पूरा होगा। ग्राहक को checkout पर अंतिम सहमत देय कीमत दिखेगी।',size=10.5)
h('क्या शामिल है')
p('प्रस्तावित सामान्य मानक A4 आकार 210 × 297 मिमी, सफेद 75 GSM साधारण कागज और पढ़ने योग्य सामान्य दस्तावेज प्रिंट है। दूसरी paper specification हो तो दुकान और PrintKarr पहले लिखित सहमति करेंगे और ग्राहक को बताएंगे। एक शीट की दोनों तरफ प्रिंट होने पर दो printed sides गिनी जाएंगी। उदाहरण के लिए 10 शीट पर दोनों तरफ B&W प्रिंट की 20 साइड का प्रिंट मूल्य ₹40 होगा।')
p('फोटो, विशेष कागज, भारी रंग कवरेज, binding, lamination और अन्य काम इन मानक दरों में शामिल नहीं हैं। प्रत्येक सेवा की दर और specification पहले तय और ग्राहक को दिखाकर मंजूर होगी। दुकान ऑर्डर स्वीकार करने के बाद छिपा surcharge नहीं जोड़ेगी।')
h('एक ही कमाई मॉडल')
p('सामान्य प्रिंट पर प्रस्तावित हिस्सेदारी 10% है। आधार सहमत प्रिंट सेवा मूल्य होगा, जिसमें अलग delivery charge और लागू वैधानिक कर शामिल नहीं होंगे; लौटाई गई राशि और दुकान की सहमति से दिए गए shop-funded discount का समायोजन होगा। अलग wholesale spread काटने के साथ दूसरा 10% commission नहीं लगेगा।')
p('Platform-funded छूट का खर्च PrintKarr उठाएगा; इससे दुकान के सहमत आधार या payout में कटौती नहीं होगी, जब तक दुकान अलग से लिखित सहमति न दे। Joining fee और मासिक fee ₹0 प्रस्तावित हैं। ऑर्डर संख्या या कमाई की गारंटी नहीं है। Stationery और add-on का हिस्सा प्रत्येक SKU या सेवा के लिए listing से पहले अलग तय होगा; कोई तय 20% margin नहीं माना जाएगा। Rider को दी जाने वाली pass-through राशि पर अतिरिक्त commission नहीं लगेगा।')

page('3 ऑर्डर स्वीकार करने से ग्राहक तक प्रक्रिया')
h('दुकान की तैयारी और ऑर्डर की स्वीकृति')
p('Onboarding में दुकान का नाम, मालिक या अधिकृत व्यक्ति, संपर्क, स्थान, खुलने के दिन व समय, मशीन, उपलब्ध सेवाएं, paper stock और दैनिक क्षमता साझा होगी। शुरुआत से पहले एक सामान्य B&W और colour sample से गुणवत्ता तथा customer specification मिलाई जाएगी।')
p('हर ऑर्डर में job ID, फाइल, चयनित pages, copies, colour, single या double side, paper और कोई सहमत add-on स्पष्ट होना चाहिए। स्वीकार करने के लिए 10 मिनट की अवधि प्रस्तावित है; अंतिम अवधि नीचे के विवरण में तय होगी। तैयारी का समय प्रत्येक ऑर्डर पर अलग स्वीकार किया जाएगा। मशीन खराब, stock कम, अस्पष्ट फाइल या समय उपलब्ध नहीं हो तो काम स्वीकार करने से पहले मना करें या स्पष्टीकरण मांगें।')
p('स्वीकृति न मिलने पर PrintKarr ग्राहक को स्थिति बताएगा और उसकी सहमति से विकल्प या refund का समन्वय करेगा। स्वीकृत काम में देरी या समस्या आते ही दुकान नामित संपर्क को बताएगी। ग्राहक की मंजूरी बिना सामग्री, साइड, कागज या कीमत नहीं बदली जाएगी। शुरुआत का समन्वय नामित संपर्क और साझा ऑर्डर रिकॉर्ड से होगा; स्वचालित दुकान आवंटन या automated settlement की सुविधा का वादा नहीं है।')
h('जिम्मेदारियां')
table(['पक्ष','सहमति के अनुसार जिम्मेदारी'],[
    ['पार्टनर दुकान','सही प्रिंट सेटिंग, गुणवत्ता, उपलब्ध stock, सुरक्षित फाइल handling, packing और स्वीकार किए गए समय पर तैयारी।'],
    ['PrintKarr','ऑर्डर और कीमत स्पष्ट करना, ग्राहक संवाद, भुगतान समन्वय, उपलब्ध rider assignment और शिकायत का एक संपर्क।'],
    ['राइडर','सहमत route, सुरक्षित transport, pickup और delivery handover तथा देरी या नुकसान की सूचना।'],
],[1.5,5.5])
h('Pickup और handover')
p('पैक पर job ID और जरूरी handover विवरण रखें; संवेदनशील दस्तावेज का नाम या प्रिंट सामग्री बाहर न दिखाएं। खुले प्रिंट के बजाय उपयुक्त folder या लिफाफे में भेजें। दुकान से rider को देते समय समय, job ID और confirmation दर्ज हो; ग्राहक handover की पुष्टि OTP, हस्ताक्षर या सहमत रिकॉर्ड से हो सकती है। प्रमाण के लिए दस्तावेज की सामग्री की फोटो न लें।')
p('डिलीवरी की उपलब्धता, क्षेत्र, fee और समय भुगतान से पहले दिखाए जाएंगे। Pickup तभी दिया जाएगा जब दुकान, स्थान और समय ग्राहक के लिए स्पष्ट व उपलब्ध हों। Hostel या office के साझा collection point को ग्राहक अलग से चुने; इसे घर के दरवाजे तक delivery नहीं कहा जाएगा। बिना वास्तविक व्यवस्था के express delivery का वादा नहीं होगा।')

page('4 भुगतान सेटलमेंट और समस्या का समाधान')
h('साप्ताहिक हिसाब और बैंक payout')
p('प्रस्ताव है कि सोमवार को पिछले सोमवार से रविवार तक पूरे हुए भुगतान वाले ऑर्डर का itemized statement दिया जाए और उसके बाद 2 कार्यदिवस में सहमत बैंक खाते में payout हो। अंतिम सप्ताह, दिन और cutoff दोनों पक्ष लिखित रूप से तय करेंगे। Statement में job ID, सेवा मूल्य, tax treatment, commission base, हिस्सेदारी, जिम्मेदार पक्ष की छूट, refund और शुद्ध देय राशि अलग दिखेगी।')
p('दुकान 3 कार्यदिवस में अंतर या आपत्ति बताए। केवल संबंधित विवादित राशि कारण व सूचना के साथ रोकना प्रस्तावित है; बाकी कमाई सामान्य समय पर दी जाए। शुरुआती समीक्षा 7 कार्यदिवस में हो। अधिक समय चाहिए तो कारण और अगली समीक्षा तारीख लिखित रूप से मिले। Earned payout अनिश्चित समय तक नहीं रोका जाएगा और कोई मनमाना penalty या छिपी कटौती नहीं होगी।')
h('Cancellation और partial refund')
p('स्वीकार करने या प्रिंट शुरू करने से पहले cancellation की सूचना तुरंत साझा करें। जो सेवा नहीं हुई और जिसकी रकम ग्राहक को वापस हुई, उस हिस्से पर commission या दुकान payout नहीं रहेगा। काम शुरू होने के बाद वास्तविक पूरे हुए काम, प्रमाण और ग्राहक के लागू अधिकार देखकर फैसला होगा; केवल cancellation के नाम पर पूरा शुल्क रोकना नियम नहीं होगा। बाद का refund अगले statement में स्पष्ट reversal के रूप में दिखेगा।')
p('सरल बिना कर वाला उदाहरण: 10 B&W साइड ₹20 और 10 colour साइड ₹50 का मूल्य ₹70 है। 10% हिस्सा ₹7 और दुकान payout ₹63 है। यदि दुकान की गलती से 5 colour साइड के ₹25 वापस हुए, बचा आधार ₹45, हिस्सा ₹4.50 और payout ₹40.50 होगा। पहले payout हो चुका हो तो अगले statement में ₹22.50 दुकान payout और ₹2.50 platform हिस्से का reversal दिखेगा। Delivery का निर्णय अलग जिम्मेदारी के अनुसार होगा।',size=10.5)
h('गलती किसकी है यह प्रमाण से तय होगा')
p('दुकान की गलत setting, खराब गुणवत्ता या छूटी साइड पर संबंधित items का reprint या refund दुकान के खर्च से होगा। ग्राहक द्वारा मंजूर फाइल की अपनी गलती अपने आप दुकान की गलती नहीं है। Chargeable reprint से पहले कारण, नई कीमत और ग्राहक की सहमति ली जाएगी।')
p('Transport damage, गलत handover या platform coordination की गलती जिम्मेदार पक्ष की अलग सहमत व्यवस्था के अनुसार संभाली जाएगी; इसे स्वतः दुकान पर नहीं डालेंगे। शिकायत PrintKarr को मिलेगी, जो दुकान और rider से सुरक्षित प्रमाण लेकर ग्राहक को कार्रवाई बताएगा। प्रस्तावित शिकायत संपर्क और समीक्षा समय नीचे भरे जाएंगे। ये नियम लागू ग्राहक अधिकारों को हटाने या blanket no-refund शर्त बनाने के लिए नहीं हैं।')

page('5 फाइल सुरक्षा ऑफर और निष्पक्ष साझेदारी')
h('फाइल और ग्राहक की जानकारी')
p('ग्राहक फाइल केवल संबंधित job के लिए साझा हो। दुकान और rider उतनी ही जानकारी देखें जितनी उनके काम के लिए जरूरी है। फाइल, नंबर या address को असंबंधित marketing, व्यक्तिगत उपयोग या बिना सहमति आगे साझा नहीं किया जाएगा। ग्राहक data से बाहर के ऑर्डर मांगने के लिए संपर्क नहीं करेंगे। दुकान के अपने मौजूदा ग्राहकों और कारोबार पर कोई व्यापक रोक नहीं होगी।')
p('फाइलों को सीमित पहुंच वाले device या folder में रखें। खुले personal chat groups, सार्वजनिक link और shared desktop पर अनावश्यक copies न बनाएं। काम पूरा होने पर print queue, temporary downloads और अनावश्यक copies समय पर हटाएं; प्रस्तावित अधिकतम अवधि काम पूरा होने के 24 घंटे बाद है, जब तक किसी खुले विवाद के लिए संबंधित सामग्री सुरक्षित रखना जरूरी न हो।')
p('हिसाब और विवाद के लिए न्यूनतम transaction record रखा जा सकता है; इसका मतलब सभी दस्तावेजों की सामग्री लंबे समय तक रखना नहीं है। Retention और जिम्मेदार संपर्क onboarding में लिखें। खोई फाइल, गलत व्यक्ति को शेयर या संदिग्ध access की सूचना तुरंत PrintKarr को दें और आगे sharing रोकें।')
h('छूट और free delivery की लागत')
p('PrintKarr का promotional discount या delivery subsidy shop payout से अपने आप नहीं कटेगा। दुकान की सहायता चाहिए तो राशि, अवधि और खर्च का मालिक पहले लिखित रूप से तय होगा। Free delivery केवल उपलब्ध retained margin, वास्तविक rider खर्च और पहले स्वीकृत platform budget पर आधारित offer हो सकती है। ₹149 का basket अपने आप margin या free delivery की गारंटी नहीं है।')
p('किसी सीमित offer का क्षेत्र, fixed slot, अवधि, अधिकतम उपयोग या कुल budget और exclusions ग्राहक को भुगतान से पहले दिखें। Print plus stationery के margin अलग जांचें। Fixed slot या collection point की eligibility हो तो उसे साफ बताएं; स्वतंत्र doorstep trip उसी कीमत पर मिलने की गारंटी न दें।')
h('बदलाव विराम और बाहर निकलना')
p('व्यावसायिक दर या प्रक्रिया में बदलाव भविष्य के ऑर्डर पर लिखित सूचना और सहमति से लागू होगा। पहले स्वीकार हुए ऑर्डर की commission दर बाद में नहीं बदलेगी। सामान्य निकास के लिए 7 दिन की सूचना प्रस्तावित है। सुरक्षा, गंभीर गुणवत्ता समस्या या धोखाधड़ी का जोखिम हो तो तुरंत नए ऑर्डर रोके जा सकते हैं, कारण और जवाब देने का अवसर दिया जाएगा।')
p('विराम या निकास के बाद भी पहले स्वीकार हुए ऑर्डर, ग्राहक refund, खुले विवाद और कमाई का अंतिम हिसाब पूरा होगा। सामान्य final statement निकास के 7 कार्यदिवस में देना प्रस्तावित है; केवल संबंधित विवादित राशि की समीक्षा अलग चलेगी। दोनों पक्ष ग्राहक को अधूरा काम या अनुत्तरित शिकायत देकर नहीं छोड़ेंगे।')

page('6 डिलीवरी का हिसाब और ग्राहक बनाने का पायलट')
p('ये स्थानीय rider quote या मुनाफे के वादे नहीं, बातचीत के गणना उदाहरण हैं। नीचे print share 10% है और tax, acquisition तथा अन्य खर्च शामिल नहीं हैं। दुकान का हिस्सा निकालने के बाद delivery के लिए बची रकम देखनी होगी।')
table(['उदाहरण','ग्राहक कुल','Platform की उपलब्ध रकम','Rider खर्च और बाकी'],[
    ['₹59 प्रिंट और अकेली doorstep delivery ₹30','₹89','₹5.90 + ₹30 = ₹35.90','अनुमान ₹60–80; कमी ₹24.10–44.10, अन्य खर्च से पहले।'],
    ['3 अलग घर एक route पर\nहर ऑर्डर ₹59 + ₹30','₹89 प्रति ऑर्डर','₹35.90 प्रति ऑर्डर','₹120 route ÷ 3 = ₹40; कमी ₹4.10 प्रति ऑर्डर।'],
    ['4 ऑर्डर एक hostel या office collection point\nहर ऑर्डर ₹59 + ₹20','₹79 प्रति ऑर्डर','₹5.90 + ₹20 = ₹25.90','₹80 route ÷ 4 = ₹20; बाकी ₹5.90 प्रति ऑर्डर।'],
    ['₹149 प्रिंट मूल्य पर delivery fee शून्य','₹149','केवल ₹14.90 प्रिंट share','Rider और अन्य खर्च अलग; मुफ्त delivery स्वतः व्यवहार्य नहीं।'],
],[2.6,1,1.55,1.85])
p('Rider का payout basket value से नहीं, पूरे समय, दूरी, waiting और उचित कमाई से तय करें। केवल बातचीत का उदाहरण: कमाई लक्ष्य ₹120 प्रति घंटा, bike खर्च ₹3 प्रति किमी, 30 मिनट और कुल 6 किमी तो ₹60 + ₹18 = ₹78 बनता है। Pickup तक जाना और उचित repositioning भी दूरी में शामिल हों। वास्तविक local agreement अलग होगा; rider को कम भुगतान करके profit दिखाना उचित नहीं है।',size=10.5)
h('पहले 30 दिन ग्राहक कैसे खोजें')
p('एक चाला micro-area और 2–3 उपयुक्त दुकानें चुनें। अनुमति लेकर एक hostel, college या office समूह तक परिचय पहुंचाएं। Counter पर QR या link से upload और pre-order pickup का workflow उपलब्ध होने पर ही प्रस्तावित करें। Fixed slots पर वास्तविक पास-पास ऑर्डर और सहमत shared collection point आजमाएं; पहले से batching होने का अनुमान न लगाएं।')
p('स्पष्ट paid delivery और उपलब्ध pickup विकल्प रखें। परिचय offer हो तो PrintKarr का लिखित budget, प्रति ऑर्डर सीमा और कुल अवधि तय हो; दुकान से बिना सहमति कटौती न हो। शुरुआत के लक्ष्य प्रयोग के होंगे, ग्राहक या order volume का वादा नहीं।')
p('हर सप्ताह outreach से paid order conversion, delivery विकल्प, वास्तविक route समय व किमी, rider cost, शिकायत और repeat orders दर्ज करें। Gross commission से rider, payment, packing, refund और promotion के variable खर्च निकालकर contribution देखें; acquisition, overhead और वास्तविक tax के बाद ही net profit देखें। विस्तार तभी करें जब दुकान margin, contribution, quality और दोबारा मांग का प्रमाण मिले। 5 किमी विस्तार बाद का निर्णय है; दमन अलग स्थानीय partner cluster हो सकता है, समान cross-city fee की गारंटी नहीं।')

page('7 दुकान की onboarding जानकारी')
p('दुकान मालिक या अधिकृत व्यक्ति यह जानकारी भरें। बैंक विवरण नामित सुरक्षित माध्यम से सत्यापित होगा; OTP, PIN या password कभी साझा न करें। यह पृष्ठ अपने आप अंतिम वाणिज्यिक सहमति नहीं बनाता।',size=10.5)
field('दुकान का नाम')
field('मालिक या अधिकृत व्यक्ति का नाम')
field('मोबाइल और email')
field('पूरा दुकान पता और landmark')
field('खुलने के दिन और समय')
field('B&W और colour मशीन तथा उपलब्ध सेवाएं')
field('Paper specification और सामान्य stock क्षमता')
field('अधिकतम स्वीकार्य काम और तैयारी समय')
field('Pickup या rider handover का नामित संपर्क')
field('शिकायत संपर्क और जवाब देने का सहमत समय')
field('फाइल deletion अवधि और जिम्मेदार व्यक्ति')
field('Bank verification संदर्भ और पूरा होने की तारीख')
p('Account holder का नाम  _______________________   बैंक का नाम  _______________________')
p('सुरक्षित चैनल या सत्यापन संपर्क  __________________________________________________')
p('Sample print check की तारीख  __________________   जांच करने वाले व्यक्ति  ______________')
p('जांच परिणाम और जरूरी सुधार  ____________________________________________________')

page('8 प्रस्तावित वाणिज्यिक विवरण और स्वीकृति')
p('दरें और समय अभी प्रस्तावित हैं। दोनों पक्ष नीचे अंतिम स्वीकार किए गए विवरण भरें, आवश्यक tax और परिचालन समीक्षा पूरी करें और हस्ताक्षर के बाद पायलट शुरू करने की तारीख तय करें। खाली या असहमत मद पर सेवा सूचीबद्ध न करें।',size=10.5)
field('PrintKarr संचालक का वास्तविक कानूनी नाम')
field('संचालक का पता और अधिकृत संपर्क')
field('पार्टनर दुकान और कानूनी मालिक का नाम')
p('पायलट आरंभ  __________________   समाप्ति  __________________   अवधि  30 दिन प्रस्तावित')
table(['सेवा','अंतिम ग्राहक दर','अंतिम दुकान payout'],[
    ['B&W प्रति प्रिंट साइड','₹ __________','₹ __________'],
    ['Colour प्रति प्रिंट साइड','₹ __________','₹ __________'],
],[2.4,2.3,2.3])
p('अंतिम commission और आधार  __________________   Joining या monthly fee  _____________')
p('सहमत paper specification तथा stationery या add-on दर सूची  __________________________')
p('Rider payout का time और km formula तथा waiting charge  ____________________________')
p('सहमत route payout या route rate list  ______________________________________________')
p('ग्राहक delivery fee और pickup fee  ________________________________________________')
p('Delivery क्षेत्र radius और उपलब्ध slots  ___________________________________________')
p('Free delivery eligibility और न्यूनतम basket  _______________________________________')
p('Promotion budget अवधि सीमा और खर्च देने वाला पक्ष  _________________________________')
p('Invoice जारी करने वाला पक्ष लागू taxes और deductions  ______________________________')
p('Statement cutoff payout दिन और dispute समीक्षा  ___________________________________')
p('Acceptance समय तथा विराम निकास और file handling में सहमत बदलाव  ____________________')
p('यहां खाली छोड़ी गई मदें अंतिम स्वीकृति नहीं हैं। पृष्ठ 2 और 6 के 10% तथा ₹2 और ₹5 दरें और rider गणनाएं केवल प्रस्ताव या उदाहरण हैं। अंतिम terms ऊपर भरकर दोनों पक्ष लिखित रूप से स्वीकार करेंगे।',size=10.5)
p('हमने दरें, commission base, जिम्मेदारियां और जोखिम पर चर्चा की है। Guaranteed order volume, blanket free delivery या बिना सहमति shop-funded discount पर सहमति नहीं मानी जाएगी। अलग दर सूची या लिखित बदलाव इसी स्वीकार किए गए विवरण के साथ रखे जाएंगे।',size=10.5)
p('दुकान प्रतिनिधि  __________________  हस्ताक्षर  ______________  तारीख  ______________')
p('PrintKarr प्रतिनिधि  _______________  हस्ताक्षर  ______________  तारीख  ______________')

# Verify arithmetic and the document's commercial boundaries before saving.
assert 59+30==89 and 59*.1+30==35.9
assert round(149*.1,2)==14.90 and round(70*.9-45*.9,2)==22.50
text='\n'.join(a.text for a in doc.paragraphs)
assert 'ग्राहक संख्या शून्य' in text and 'कानूनी नाम' in text
doc.core_properties.title='PrintKarr दुकान साझेदारी प्रस्ताव'
doc.core_properties.subject='स्थानीय प्रिंट दुकानों के लिए प्रस्तावित 30 दिन का पायलट और वाणिज्यिक विवरण'
doc.core_properties.author='PrintKarr'
doc.core_properties.keywords='Hindi shop partner pilot proposal'
doc.save(OUT)
print(OUT)
